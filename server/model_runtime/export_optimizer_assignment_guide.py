from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/home/ubuntu/exports/CarbonSense_Optimizer_Recommendation_Assignment_Guide.docx")

FOREST = "1B4D3E"
MINT = "D4F0E5"
PALE = "F4F9F7"
CHARCOAL = "2D3748"
AMBER = "F59E0B"


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def set_cell_text(cell, text: str, bold: bool = False, color: str = CHARCOAL, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor.from_string(FOREST)
    run.font.size = Pt(16 if level == 1 else 12)


def add_body(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        prefix.font.name = "Aptos"
        prefix.font.size = Pt(10.5)
        remainder = paragraph.add_run(text[len(bold_prefix):])
        remainder.font.name = "Aptos"
        remainder.font.size = Pt(10.5)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_callout(document: Document, title: str, text: str, fill: str = MINT) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.width = Inches(6.65)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(4)
    title_run = paragraph.add_run(title + "\n")
    title_run.bold = True
    title_run.font.name = "Aptos"
    title_run.font.size = Pt(10.5)
    title_run.font.color.rgb = RGBColor.from_string(FOREST)
    text_run = paragraph.add_run(text)
    text_run.font.name = "Aptos"
    text_run.font.size = Pt(10)
    text_run.font.color.rgb = RGBColor.from_string(CHARCOAL)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_step_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.78, 2.12, 3.72]
    for cell, title, width in zip(table.rows[0].cells, ["Step", "What the optimizer does", "What this means for the user"], widths):
        cell.width = Inches(width)
        shade(cell, FOREST)
        set_cell_text(cell, title, bold=True, color="FFFFFF", size=9)
    rows = [
        ("01", "Read the selected start and target", "For example: AI estimate 715 kg/month with a 28% target requires roughly 200.2 kg/month of planned saving."),
        ("02", "Create eligible actions from the saved profile", "The action set can include high-impact meals, travel distance, clothing, electricity reference activity, and landfill waste."),
        ("03", "Calculate a bounded potential for each action", "Each action has a supported maximum change and an estimated saving per unit. The optimizer cannot exceed that bound."),
        ("04", "Select the lowest-effort combination", "The planner allocates only the amount needed from eligible actions until the chosen target is met or the action bounds are exhausted."),
        ("05", "Rank active and optional actions", "In-plan actions appear first by actual planned saving. Unused options follow as optional capacity for a higher target or a different mix of actions."),
    ]
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            shade(cell, "FFFFFF" if index % 2 == 0 else PALE)
        set_cell_text(cells[0], row[0], bold=True, color=FOREST)
        set_cell_text(cells[1], row[1], bold=True)
        set_cell_text(cells[2], row[2])


def add_example_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headings = ["Displayed item", "Value", "Why it appears", "Interpretation"]
    widths = [1.7, 1.15, 1.95, 1.82]
    for cell, title, width in zip(table.rows[0].cells, headings, widths):
        cell.width = Inches(width)
        shade(cell, FOREST)
        set_cell_text(cell, title, bold=True, color="FFFFFF", size=8.5)
    rows = [
        ("Priority 01 · In plan", "−200.2 kg/month", "Its available saving can meet the 28% target within the configured bound.", "It is the active action chosen for this particular target."),
        ("Priority 02 · Optional", "Up to −48.5 kg/month", "Travel is not needed after the first action reaches the target.", "It is extra capacity for a higher target or a more diversified plan."),
        ("Priority 03 onward", "Supported bounds", "Other profile-relevant actions are still feasible but unused.", "They are alternatives, not compulsory tasks."),
    ]
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            shade(cell, "FFFFFF" if index % 2 == 0 else PALE)
        for cell, text in zip(cells, row):
            set_cell_text(cell, text, bold=(cell is cells[0]))


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("How CarbonSense Optimizer Recommendations Are Assigned")
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(23)
    title_run.font.color.rgb = RGBColor.from_string(FOREST)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("Plain-language guide for the result-led reduction planner")
    subtitle_run.italic = True
    subtitle_run.font.name = "Aptos"
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor.from_string("52635D")

    add_callout(
        document,
        "Short answer",
        "Optimizer recommendations are generated automatically from the saved questionnaire profile, the selected starting method, and the chosen reduction target. They are not assigned by another person and are not a second trained machine-learning model.",
    )

    add_heading(document, "1. The recommendation-assignment sequence")
    add_body(document, "The optimizer begins after a user has completed both an AI prediction and a transparent baseline. The user chooses one of those results as the planning start, then chooses a percentage reduction target. CarbonSense uses only the selected result for that target; it does not average the AI estimate and baseline.")
    add_step_table(document)

    add_heading(document, "2. Worked example: a 28% target")
    add_body(document, "Suppose the selected AI estimate is 715 kgCO₂e/month. A 28% reduction target corresponds to about 200.2 kgCO₂e/month of planned saving. If the high-impact-meal action has enough available capacity to supply that full saving, the optimizer assigns it as Priority 01 and labels it In plan.")
    add_example_table(document)

    add_callout(
        document,
        "Why one action can appear first",
        "The planner may not need multiple active actions at a moderate target. This does not mean the other actions are irrelevant; it means the selected first action already meets the requested saving within its configured limit. At a higher target, the optimizer can activate further actions in sequence.",
        fill="FFF4D8",
    )

    add_heading(document, "3. What the labels mean")
    add_body(document, "In plan means the action receives a calculated share of the current target. Optional means the action is compatible with the profile but is not currently required; it is shown as extra capacity. Priority order is not a moral judgement and does not prove an action will achieve the displayed saving in real life. It is a transparent planning order within the optimizer’s saved assumptions and limits.")

    add_heading(document, "4. Optimizer recommendations versus organization assignments")
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, title in zip(table.rows[0].cells, ["Type", "Who creates it", "Purpose"]):
        shade(cell, FOREST)
        set_cell_text(cell, title, bold=True, color="FFFFFF", size=9)
    for row in [
        ("Optimizer recommendation", "CarbonSense automatically", "Builds a profile-specific plan from the selected result and target."),
        ("Organization assignment", "An authorized organization administrator", "Sends a separate approved catalog action to a member for organizational follow-up."),
    ]:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            shade(cells[index], PALE if row[0].startswith("Optimizer") else "FFFFFF")
            set_cell_text(cells[index], text, bold=(index == 0))

    add_heading(document, "5. Safe interpretation")
    add_body(document, "Treat the recommendation list as an indicative action-planning aid. The displayed values are modelled savings based on the saved profile, bounded action settings, and documented factor or planning assumptions. They are not direct emissions measurements, verified personal reductions, or a guarantee that a specific lifestyle change will produce the exact displayed amount.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("CarbonSense · Optimizer recommendation assignment guide · Indicative planning aid")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("60716A")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
