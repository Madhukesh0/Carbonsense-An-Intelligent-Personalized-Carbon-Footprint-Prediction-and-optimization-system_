"""Export the deterministic XGBoost QA audit as a clearly labeled Word report."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
AUDIT_PATH = ROOT / "audits" / "xgb_deterministic_100_scenario_audit.json"
DEFAULT_OUTPUT = Path("/home/ubuntu/exports/CarbonSense_XGBoost_100_Scenario_QA.docx")

GREEN = RGBColor(21, 127, 84)
DARK = RGBColor(16, 42, 45)
AMBER = "FFF4D6"


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 8, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = DARK if level == 1 else GREEN


def format_drivers(drivers: list[dict]) -> str:
    return "; ".join(f"{driver['label']} ({driver['direction']} {driver['shapValue']:+.1f})" for driver in drivers)


def input_summary(payload: dict) -> str:
    return (
        f"Diet: {payload['diet']}; transport: {payload['transport']}; distance: {payload['vehicle_monthly_distance_km']} km; "
        f"air: {payload['frequency_of_traveling_by_air']}; grid: {payload['region']}; grocery: {payload['monthly_grocery_bill']}; "
        f"clothes: {payload['how_many_new_clothes_monthly']}; TV/PC: {payload['how_long_tv_pc_daily_hour']} h; internet: {payload['how_long_internet_daily_hour']} h"
    )


def configure_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def build_document(audit: dict) -> Document:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10)
    document.core_properties.title = "CarbonSense XGBoost 100-Scenario QA Report"
    document.core_properties.author = "Manus AI"

    title = document.add_heading("CarbonSense XGBoost 100-Scenario QA Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = DARK
    subtitle = document.add_paragraph("Frozen xgb-final-54f-v2 · Deterministic simulated test evidence")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.italic = True
        run.font.color.rgb = GREEN

    notice = document.add_table(rows=1, cols=1)
    notice_cell = notice.cell(0, 0)
    shade_cell(notice_cell, AMBER)
    set_cell_text(
        notice_cell,
        "IMPORTANT: This document contains deterministic simulated QA scenarios. The 100 rows are not real people, user records, training data, or a population emissions distribution.",
        bold=True,
        size=10,
        color=DARK,
    )

    add_heading(document, "Purpose and Validation Boundary", 1)
    document.add_paragraph(
        "This report validates the same persisted XGBoost runtime used by CarbonSense. It exercises bounded low-impact, typical, and high-impact inputs to check that predictions are finite, the 34-to-54 feature transformation remains intact, and ranked live contribution values reconcile to the prediction. It does not validate real-world personal emissions or establish causal environmental impacts."
    )

    add_heading(document, "QA Summary", 1)
    summary = audit["summary"]
    summary_table = document.add_table(rows=1, cols=2)
    summary_table.style = "Light Shading Accent 1"
    set_cell_text(summary_table.cell(0, 0), "Measure", bold=True, color=RGBColor(255, 255, 255))
    set_cell_text(summary_table.cell(0, 1), "Result (kgCO₂e/month)", bold=True, color=RGBColor(255, 255, 255))
    shade_cell(summary_table.cell(0, 0), "157F54")
    shade_cell(summary_table.cell(0, 1), "157F54")
    for label, value in [
        ("Lowest deterministic test result", summary["minimumKg"]),
        ("Highest deterministic test result", summary["maximumKg"]),
        ("Matrix mean", summary["meanKg"]),
        ("Matrix median", summary["medianKg"]),
    ]:
        cells = summary_table.add_row().cells
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], f"{value:.1f}")

    add_heading(document, "Profile Bands", 1)
    band_table = document.add_table(rows=1, cols=4)
    band_table.style = "Light Shading Accent 1"
    for index, header in enumerate(["Test band", "Scenarios", "Mean", "Minimum–maximum"]):
        set_cell_text(band_table.cell(0, index), header, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(band_table.cell(0, index), "157F54")
    labels = {"low_impact": "Low impact", "typical": "Typical", "high_impact": "High impact"}
    for profile, values in summary["byProfile"].items():
        cells = band_table.add_row().cells
        set_cell_text(cells[0], labels[profile])
        set_cell_text(cells[1], values["count"])
        set_cell_text(cells[2], f"{values['meanKg']:.1f}")
        set_cell_text(cells[3], f"{values['minimumKg']:.1f}–{values['maximumKg']:.1f}")

    add_heading(document, "Representative Live Model Explanations", 1)
    rep_table = document.add_table(rows=1, cols=3)
    rep_table.style = "Light Shading Accent 1"
    for index, header in enumerate(["Case", "Prediction", "Top live XGBoost drivers"]):
        set_cell_text(rep_table.cell(0, index), header, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(rep_table.cell(0, index), "157F54")
    case_labels = {"lowestPrediction": "Lowest test result", "closestToMatrixMean": "Closest to matrix mean", "highestPrediction": "Highest test result"}
    for key, label in case_labels.items():
        row = audit["representativeCases"][key]
        cells = rep_table.add_row().cells
        set_cell_text(cells[0], label)
        set_cell_text(cells[1], f"{row['predictionKg']:.1f} kgCO₂e/month")
        set_cell_text(cells[2], format_drivers(row["topDrivers"]), size=8)

    document.add_page_break()
    add_heading(document, "Full Deterministic Scenario Evidence", 1)
    paragraph = document.add_paragraph(
        "The following 100 rows are the full simulated QA matrix. Each row was passed through the frozen live XGBoost runtime. The reconciliation delta is the absolute difference between grouped contribution reconstruction and the prediction; all rows are within 0.2 kgCO₂e."
    )
    for run in paragraph.runs:
        run.font.size = Pt(9)

    section = document.sections[-1]
    configure_landscape(section)
    columns = ["ID", "Band", "Selected bounded test inputs", "Prediction", "Range", "Recon. Δ", "Top live drivers"]
    scenario_table = document.add_table(rows=1, cols=len(columns))
    scenario_table.style = "Table Grid"
    for index, header in enumerate(columns):
        set_cell_text(scenario_table.cell(0, index), header, bold=True, size=7, color=RGBColor(255, 255, 255))
        shade_cell(scenario_table.cell(0, index), "157F54")
    for scenario in audit["scenarios"]:
        cells = scenario_table.add_row().cells
        values = [
            scenario["scenarioId"],
            labels[scenario["profile"]],
            input_summary(scenario["input"]),
            f"{scenario['predictionKg']:.1f}",
            f"{scenario['uncertaintyRange']['low']:.1f}–{scenario['uncertaintyRange']['high']:.1f}",
            f"{scenario['reconciliationDelta']:.3f}",
            format_drivers(scenario["topDrivers"]),
        ]
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, size=6.5)
    widths = [0.65, 0.75, 3.2, 0.7, 0.9, 0.7, 2.25]
    for row in scenario_table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)

    document.add_paragraph()
    boundary = document.add_paragraph()
    boundary.add_run("Interpretation boundary: ").bold = True
    boundary.add_run("These outcomes are deterministic test evidence for the deployed model. They describe model behavior under simulated inputs; they are not direct emissions measurements, causal estimates, or real-person outcomes.")
    return document


def main() -> None:
    if not AUDIT_PATH.exists():
        raise FileNotFoundError(f"Missing deterministic audit: {AUDIT_PATH}")
    audit = json.loads(AUDIT_PATH.read_text())
    document = build_document(audit)
    DEFAULT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(DEFAULT_OUTPUT)
    print(json.dumps({"output": str(DEFAULT_OUTPUT), "scenarioCount": audit["scenarioCount"]}))


if __name__ == "__main__":
    main()
